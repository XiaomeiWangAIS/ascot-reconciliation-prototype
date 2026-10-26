"""Contract checks for SUN allocation preservation and schema normalization."""
import csv
import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from src.ascot.sanitize_sun_jnl import sanitize_workbook


def test_reordered_aliases_optional_fields_and_duplicates(tmp_path):
    path = tmp_path / 'Paid Invoices for SUN JNL July 2026.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Page1_1'
    sheet.append(['Report ran at Aug 3, 2026'])
    sheet.append(['Paid Invoices for SUN JNL'])
    sheet.append(['Paid Date Between Jul 1, 2026 and Jul 31, 2026'])
    sheet.append([])
    sheet.append([' Paid Date ', ' Description ', 'Amount', ' Account ',
                  'Bank Account Name (Custom 9)', 'VAT Amount ', 'Total Amount',
                  'Trans Ref', 'Payment Type', 'Company'])
    row = [datetime(2026, 7, 2), 'Vendor', -12.34, '00102', 'Bank',
           1.01, 90, '000123', '04447', 'AAIC']
    sheet.append(row)
    sheet.append(row)
    other = book.create_sheet('AIC')
    other.append(['This must never be ingested'])
    book.save(path)
    records, qa = sanitize_workbook(path)
    assert len(records) == qa['input_detail_row_count'] == 2
    assert qa['exact_duplicate_row_count'] == 1
    assert qa['header_row'] == 5
    assert qa['aic_row_count'] == qa['aicre_row_count'] == 0
    assert records[0]['source_row'] == 6
    assert records[0]['amount'] == Decimal('-12.34')
    assert records[0]['vat_amount'] == Decimal('1.01')
    assert records[0]['total_amount'] == Decimal('90')
    assert records[0]['request_total'] is None
    assert records[0]['account_code'] == '00102'
    assert records[0]['transaction_reference'] == '000123'
    assert records[0]['payment_type_raw'] == '04447'
    assert records[0]['report_period_end'] == '2026-07-31'
    assert json.loads(records[0]['raw_values_json'])[7] == '000123'


@pytest.mark.parametrize('bad', ['=1+2', 'invalid money'])
def test_rejects_unusable_amount(tmp_path, bad):
    path = tmp_path / 'Paid Invoices for SUN JNL January 2026.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Page1_1'
    sheet.append(['Account', 'Description', 'Amount', 'Paid Date'])
    sheet.append(['100', 'Vendor', bad, datetime(2026, 1, 7)])
    book.save(path)
    with pytest.raises(ValueError):
        sanitize_workbook(path)


@pytest.mark.parametrize('month,count,negative,zero,missing_account,missing_payment,duplicates', [
    ('January', 731, 5, 0, 0, 1, 0),
    ('February', 675, 9, 1, 0, 1, 0),
    ('Mar', 768, 39, 0, 0, 1, 1),
    ('Apr', 919, 40, 0, 0, 57, 2),
    ('May', 809, 63, 1, 0, 1, 1),
    ('June', 1244, 90, 0, 6, 1, 0),
    ('July', 1130, 6, 0, 0, 1, 0),
])
def test_observed_months(month, count, negative, zero, missing_account, missing_payment, duplicates):
    path = Path(f'data/raw/Paid Invoices for SUN JNL {month} 2026.xlsx')
    if not path.exists():
        pytest.skip('Local raw workbook not available')
    records, qa = sanitize_workbook(path)
    assert len(records) == qa['input_detail_row_count'] == count
    assert qa['negative_amount_count'] == negative
    assert qa['zero_amount_count'] == zero
    assert qa['missing_account_count'] == missing_account
    assert qa['missing_payment_type_count'] == missing_payment
    assert qa['exact_duplicate_row_count'] == duplicates
    assert all(isinstance(r['amount'], Decimal) for r in records)
    assert sum(qa['company_counts'].values()) == count
