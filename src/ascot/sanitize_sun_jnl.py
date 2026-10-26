"""Normalize SUN JNL Page1_1 allocation lines without filtering or aggregation.

Run: python -m src.ascot.sanitize_sun_jnl
Money is Decimal in the API and decimal text in CSV. Raw cell values and the
source SHA256 are retained for audit; Excel formatting has no accounting role.
"""
import argparse
import calendar
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

SOURCE_SHEET = 'Page1_1'
ALIASES = {
    'layout identifier': 'layout_identifier', 'account': 'account_code',
    'description': 'description', 'ccy': 'currency', 'amount': 'amount',
    'd/c': 'debit_credit', 'paid date': 'paid_date',
    'trans date': 'transaction_date', 'trans ref': 'transaction_reference',
    'company': 'company', 'bank account name': 'bank_account_name',
    'bank account code': 'bank_account_code', 'payment type': 'payment_type_raw',
    't1 trust fund': 'trust_fund', 't3 cost center': 'cost_center',
    't4 inter co': 'intercompany', 'loc': 'location',
    'line description': 'line_description', 'request total': 'request_total',
    'tax amount': 'tax_amount', 'payment amount': 'payment_amount',
    'vat amount': 'vat_amount', 'total amount': 'total_amount',
    'period': 'period', 't0 yoa': 'year_of_account', 't2 class': 'class',
    't5 vat': 'vat_code', 't6 lca': 'lca', 't8 emp': 'employee',
}
MONEY = {'amount', 'request_total', 'tax_amount', 'payment_amount', 'vat_amount', 'total_amount'}
DATES = {'paid_date', 'transaction_date'}
FIELDS = ['source_file', 'source_sheet', 'source_row', 'accounting_period',
          'report_run_date', 'report_period_start', 'report_period_end',
          *ALIASES.values(), 'raw_values_json']


def header_key(value):
    return re.sub(r'\s+', ' ', re.sub(r'\s*\(custom\s+9\)\s*$', '', str(value or ''), flags=re.I)).strip().lower()


def accounting_period(path):
    match = re.fullmatch(r'Paid Invoices for SUN JNL ([A-Za-z]+) (\d{4})', Path(path).stem, re.I)
    if not match:
        raise ValueError(f'Cannot identify accounting period: {path}')
    for fmt in ('%B %Y', '%b %Y'):
        try:
            return datetime.strptime(' '.join(match.groups()), fmt).strftime('%Y-%m')
        except ValueError:
            pass
    raise ValueError(f'Invalid month/year: {path}')


def json_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise TypeError(f'Unsupported raw value: {type(value)}')


def normalize(value, field):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if field in MONEY:
        try:
            result = Decimal(str(value).strip())
            if not result.is_finite():
                raise ValueError('Nonfinite amount')
            return result
        except (InvalidOperation, ValueError) as exc:
            raise ValueError(f'Invalid {field}: {value!r}') from exc
    if field in DATES:
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        # Do not guess ambiguous dates or interpret formatting as logic.
        return date.fromisoformat(str(value).strip()).isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sanitize_workbook(path):
    """Return (records, QA). Fail on ambiguous headers, formulas or invalid types."""
    path = Path(path)
    period = accounting_period(path)
    records, raw_signatures = [], Counter()
    warnings, preamble, mapping = [], [], None
    metadata = dict.fromkeys(['report_run_date', 'report_period_start', 'report_period_end'])
    with path.open('rb') as stream:
        digest = hashlib.file_digest(stream, 'sha256').hexdigest()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f'{path.name}: missing {SOURCE_SHEET}')
        sheet = workbook[SOURCE_SHEET]
        for number, cells in enumerate(sheet.iter_rows(), 1):
            values = [c.value for c in cells]
            if not any(v is not None for v in values):
                continue
            if mapping is None:
                keys = [header_key(v) for v in values]
                if not {'account', 'description', 'amount', 'paid date'}.issubset(keys):
                    preamble.extend(str(v) for v in values if v is not None)
                    continue
                header_row = number
                headers = values
                mapping = {i: ALIASES[k] for i, k in enumerate(keys) if k in ALIASES}
                if len(set(mapping.values())) != len(mapping):
                    raise ValueError(f'{path.name}: duplicate canonical headers')
                unknown = [v for v, k in zip(values, keys) if v is not None and k not in ALIASES]
                if unknown:
                    warnings.append(f'Unmapped columns retained in raw values: {unknown}')
                missing = sorted(set(ALIASES.values()) - set(mapping.values()))
                for text in preamble:
                    if text.startswith('Report ran at '):
                        metadata['report_run_date'] = datetime.strptime(text[14:], '%b %d, %Y').date().isoformat()
                    match = re.fullmatch(r'Paid Date Between (.+) and (.+)', text)
                    if match:
                        metadata['report_period_start'], metadata['report_period_end'] = [datetime.strptime(v, '%b %d, %Y').date().isoformat() for v in match.groups()]
                continue
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ValueError(f'{path.name}:{number}: formula or Excel error in canonical source')
            raw = json.dumps(values, default=json_value, ensure_ascii=False)
            raw_signatures[raw] += 1
            record = dict.fromkeys(FIELDS)
            record.update(source_file=str(path), source_sheet=SOURCE_SHEET,
                          source_row=number, accounting_period=period,
                          raw_values_json=raw, **metadata)
            for i, field in mapping.items():
                try:
                    record[field] = normalize(values[i], field)
                except ValueError as exc:
                    raise ValueError(f'{path.name}:{number}: {exc}') from exc
            records.append(record)
        if mapping is None:
            raise ValueError(f'{path.name}: transaction header not found')
    finally:
        workbook.close()
    companies = Counter(r['company'] for r in records)
    year, month = map(int, period.split('-'))
    month_end = date(year, month, calendar.monthrange(year, month)[1]).isoformat()
    if metadata['report_period_end'] and metadata['report_period_end'] < month_end:
        warnings.append('Report period ends before accounting month end')
    if metadata['report_run_date'] and metadata['report_run_date'] < month_end:
        warnings.append('Report generated before accounting month end; completeness unconfirmed')
    if not metadata['report_period_start']:
        warnings.append('Report-period metadata absent; accounting period comes from filename')
    qa = dict(accounting_period=period, source_file=str(path), source_sheet=SOURCE_SHEET,
              source_sha256=digest, ignored_sheets=[s for s in workbook.sheetnames if s != SOURCE_SHEET],
              header_row=header_row, source_headers=headers, header_mapping={str(i+1): f for i, f in mapping.items()},
              absent_optional_fields=missing, **metadata,
              input_detail_row_count=sum(raw_signatures.values()), sanitized_row_count=len(records),
              company_counts=dict(sorted(companies.items(), key=lambda x: str(x[0]))),
              aic_row_count=companies['AIC'], aicre_row_count=companies['AICRE'],
              negative_amount_count=sum(r['amount'] is not None and r['amount'] < 0 for r in records),
              zero_amount_count=sum(r['amount'] == 0 for r in records),
              missing_account_count=sum(r['account_code'] is None for r in records),
              missing_payment_type_count=sum(r['payment_type_raw'] is None for r in records),
              exact_duplicate_row_count=sum(n-1 for n in raw_signatures.values()),
              missing_amount_count=sum(r['amount'] is None for r in records),
              structural_warnings=warnings)
    for field, label in [('missing_account_count', 'Missing account codes'),
                         ('missing_payment_type_count', 'Missing payment types'),
                         ('missing_amount_count', 'Missing amounts'),
                         ('exact_duplicate_row_count', 'Exact duplicate rows retained')]:
        if qa[field]:
            warnings.append(f'{label}: {qa[field]}')
    return records, qa


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input-dir', type=Path, default=Path('data/raw'))
    parser.add_argument('--output-dir', type=Path, default=Path('data/processed/sun_jnl'))
    args = parser.parse_args()
    files = sorted(args.input_dir.glob('Paid Invoices for SUN JNL *.xlsx'), key=accounting_period)
    if not files:
        raise FileNotFoundError(f'No SUN JNL workbooks in {args.input_dir}')
    periods = [accounting_period(p) for p in files]
    if len(periods) != len(set(periods)):
        raise ValueError('Multiple source workbooks for one accounting period')
    # Validate all inputs before publishing this batch's outputs.
    results = [sanitize_workbook(p) for p in files]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for records, qa in results:
        output = args.output_dir / f"sun_jnl_{qa['accounting_period']}.csv"
        with output.open('w', newline='', encoding='utf-8') as stream:
            writer = csv.DictWriter(stream, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(records)
        qa['output_file'] = str(output)
        manifest.append(qa)
        print(f"{qa['accounting_period']}: {qa['input_detail_row_count']} input / {len(records)} sanitized")
    (args.output_dir / 'sun_jnl_manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')


if __name__ == '__main__':
    main()
