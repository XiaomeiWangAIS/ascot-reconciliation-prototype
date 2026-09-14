# ============================================================
# ASCOT AIC CASH EXPENSE AUTOMATION
# AIC WF SANITIZER
# STREAMING + CHRONOLOGICAL + EMPTYCELL-SAFE VERSION
# ============================================================
#
# THIS VERSION INCORPORATES THE DESIGN RULES WE HAVE LOCKED IN:
#
# 1. Monthly Resource Pro files are processed chronologically:
#
#       January -> February -> March -> ...
#
#    NOT alphabetically.
#
# 2. Resource Pro workbooks are read in streaming mode because
#    their internal Excel formatting makes them very large.
#
# 3. Ascot's cell colors are preserved as legacy/presentation
#    metadata, but colors are NOT treated as the authoritative
#    machine-readable accounting category.
#
# 4. The AIC WF structure is explicitly represented as:
#
#       main cash activity
#              ↓
#       subtotal/control row
#              ↓
#       presentation gap
#              ↓
#       ZBA / sweep section
#
# 5. We do NOT infer:
#
#       "large blank gap = ZBA"
#
#    Instead, we use:
#
#       subtotal boundary
#              +
#       ZBA/sweep transaction description
#
# 6. Subtotal rows are preserved separately as control rows.
#
# 7. Every normalized transaction preserves:
#
#       accounting period
#       source workbook
#       source sheet
#       source row
#
# 8. IMPORTANT TECHNICAL FIX:
#
#    In openpyxl read-only mode, empty cells may be EmptyCell
#    objects. EmptyCell does NOT have attributes such as:
#
#       .row
#       .column
#
#    Therefore this script derives row/column numbers from
#    enumerate() rather than from cell objects.
#
# ============================================================


# ============================================================
# IMPORTS
# ============================================================

import csv
import json
import re

from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook


# ============================================================
# PROJECT FOLDERS
# ============================================================

RAW_FOLDER = Path(
    "data/raw"
)

PROCESSED_FOLDER = Path(
    "data/processed"
)

PROCESSED_FOLDER.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# HOW MANY EXCEL COLUMNS TO READ
# ============================================================
#
# Resource Pro sometimes has formatting extending thousands
# of columns to the right.
#
# The meaningful AIC WF data are in the beginning of the
# worksheet, so we intentionally inspect only the first 20.
# ============================================================

MAX_COLUMNS_TO_READ = 20


# ============================================================
# CANONICAL AIC WF COLUMN NAMES
# ============================================================

HEADER_MAP = {

    "as-of date":
        "as_of_date",

    "acct no":
        "account_number",

    "acct name":
        "account_name",

    "tran desc":
        "transaction_description",

    "debit amt":
        "debit_amount",

    "credit amt":
        "credit_amount",

    "0 day flt amt":
        "zero_day_float_amount",

    "check number":
        "check_number",

    "descriptive text 1":
        "descriptive_text_1",

    "account":
        "gl_account",

    "cost center":
        "cost_center",

    "description":
        "accounting_description",

    "back up":
        "backup",
}


# ============================================================
# ASCOT / RESOURCE PRO LEGEND LABELS
# ============================================================
#
# These are legacy Excel presentation labels.
#
# We preserve them, but we do NOT automatically treat them as
# the machine's final classification.
# ============================================================

KNOWN_LEGEND_LABELS = {

    "ims":
        "IMS",

    "bank items":
        "Bank Items",

    "prior month":
        "Prior Month",

    "state payments":
        "State Payments",

    "reinsurance":
        "Reinsurance",

    "interco transfers":
        "Interco Transfers",

    "investment":
        "Investment",

    "loss funding/tpa":
        "Loss Funding/TPA",
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(value):
    """
    Convert Excel text into a normalized lowercase string.

    Example:

        "   Tran Desc   "

    becomes:

        "tran desc"
    """

    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip(),
    ).lower()


# ============================================================
# NUMBER NORMALIZATION
# ============================================================

def normalize_number(value):
    """
    Convert an Excel numeric value into a Python float.

    Empty cells remain None.

    Debit and credit are intentionally preserved separately.
    """

    if value is None:
        return None

    if isinstance(
        value,
        (int, float),
    ):

        return float(
            value
        )

    text = str(
        value
    ).strip()

    if text == "":
        return None

    text = (
        text
        .replace("$", "")
        .replace(",", "")
    )

    try:

        return float(
            text
        )

    except ValueError:

        return None


# ============================================================
# DATE NORMALIZATION
# ============================================================

def normalize_date(value):
    """
    Convert dates to YYYY-MM-DD when possible.
    """

    if value is None:
        return None

    if isinstance(
        value,
        datetime,
    ):

        return value.date().isoformat()

    if hasattr(
        value,
        "isoformat",
    ):

        try:

            return value.isoformat()

        except Exception:

            pass

    # Do not aggressively guess unusual date formats.
    return str(
        value
    ).strip()


# ============================================================
# ACCOUNTING PERIOD FROM FILENAME
# ============================================================

def extract_accounting_period(file_path):
    """
    Extract month and year from a filename such as:

        January 2026 Resource Pro Cash.xlsx

    and convert them to:

        datetime(2026, 1, 1)

    This allows true chronological sorting.
    """

    match = re.match(

        r"^([A-Za-z]+)\s+(\d{4})\b",

        file_path.stem,
    )

    if not match:

        raise ValueError(

            "Cannot identify accounting period "
            f"from filename: {file_path.name}"
        )

    month_name = (
        match.group(1).title()
    )

    year = (
        match.group(2)
    )

    month_year = (
        f"{month_name} {year}"
    )

    try:

        return datetime.strptime(
            month_year,
            "%B %Y",
        )

    except ValueError as error:

        raise ValueError(

            "Invalid month/year at beginning of filename: "
            f"{file_path.name}"
        ) from error


# ============================================================
# ACCOUNTING PERIOD LABEL
# ============================================================

def accounting_period_label(file_path):
    """
    Convert accounting period into:

        2026-01
        2026-02
        ...
    """

    period = extract_accounting_period(
        file_path
    )

    return period.strftime(
        "%Y-%m"
    )


# ============================================================
# READ CELL FILL
# ============================================================

def get_fill_key(cell):
    """
    Preserve Excel fill information.

    This function is deliberately EmptyCell-safe.

    getattr() returns None when the streamed cell does not
    expose normal style attributes.
    """

    fill = getattr(
        cell,
        "fill",
        None,
    )

    if fill is None:
        return None

    fill_type = getattr(
        fill,
        "fill_type",
        None,
    )

    if fill_type is None:
        return None

    color = getattr(
        fill,
        "fgColor",
        None,
    )

    if color is None:
        return None

    color_type = getattr(
        color,
        "type",
        None,
    )

    if color_type == "rgb":

        color_value = getattr(
            color,
            "rgb",
            None,
        )

    elif color_type == "indexed":

        color_value = getattr(
            color,
            "indexed",
            None,
        )

    elif color_type == "theme":

        color_value = getattr(
            color,
            "theme",
            None,
        )

    else:

        color_value = None

    return {

        "fill_type":
            fill_type,

        "color_type":
            color_type,

        "color_value":
            color_value,
    }


# ============================================================
# CREATE COMPARABLE COLOR SIGNATURE
# ============================================================

def fill_signature(fill_key):
    """
    Convert fill metadata into a comparable string.
    """

    if fill_key is None:
        return None

    return (
        f"{fill_key['fill_type']}|"
        f"{fill_key['color_type']}|"
        f"{fill_key['color_value']}"
    )


# ============================================================
# IDENTIFY ZBA / SWEEP ACTIVITY
# ============================================================

def is_zba_or_sweep(
    transaction_description,
    descriptive_text,
):
    """
    Identify ZBA/sweep transactions from descriptive evidence.

    We do NOT classify something as a sweep merely because
    it appears after a blank gap.
    """

    combined_text = " ".join(

        [

            str(
                transaction_description
                or ""
            ),

            str(
                descriptive_text
                or ""
            ),
        ]
    ).lower()

    keywords = [

        "zba",

        "sweep",

        "sweep principal",

        "investment sweep",
    ]

    return any(

        keyword in combined_text

        for keyword
        in keywords
    )


# ============================================================
# DETECT SUM / SUBTOTAL ROW
# ============================================================

def detect_sum_formula(row):
    """
    Search a streamed worksheet row for SUM formulas.

    Uses getattr() so EmptyCell objects are safe.
    """

    formulas = []

    for cell in row:

        data_type = getattr(
            cell,
            "data_type",
            None,
        )

        if data_type != "f":
            continue

        formula = str(
            getattr(
                cell,
                "value",
                "",
            )
        )

        formulas.append(
            formula
        )

    contains_sum = any(

        "sum(" in formula.lower()

        for formula
        in formulas
    )

    return (
        contains_sum,
        formulas,
    )


# ============================================================
# WRITE CSV
# ============================================================

def write_csv(
    path,
    records,
):
    """
    Write dictionaries to CSV using Python's built-in csv
    library so we do not need pandas.
    """

    if not records:

        # Overwrite/create a truly empty output.
        with path.open(
            "w",
            encoding="utf-8-sig",
        ):
            pass

        return

    fieldnames = []

    for record in records:

        for key in record.keys():

            if key not in fieldnames:

                fieldnames.append(
                    key
                )

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as file:

        writer = csv.DictWriter(

            file,

            fieldnames=fieldnames,
        )

        writer.writeheader()

        writer.writerows(
            records
        )


# ============================================================
# SANITIZE ONE RESOURCE PRO WORKBOOK
# ============================================================

def sanitize_workbook(
    workbook_path,
    processing_sequence,
):
    """
    Sanitize one month's AIC WF worksheet.
    """

    period = accounting_period_label(
        workbook_path
    )

    print(
        "\n========================================"
    )

    print(
        f"Processing #{processing_sequence}: "
        f"{workbook_path.name}"
    )

    print(
        f"Accounting period: {period}"
    )

    print(
        "========================================"
    )


    # ========================================================
    # LOAD WORKBOOK IN STREAMING MODE
    # ========================================================

    wb = load_workbook(

        workbook_path,

        read_only=True,

        data_only=False,

        keep_links=False,
    )


    try:

        # ====================================================
        # FIND AIC WF
        # ====================================================

        if "AIC WF" not in wb.sheetnames:

            raise ValueError(

                f"{workbook_path.name} does not contain "
                "'AIC WF'."
            )

        ws = wb[
            "AIC WF"
        ]


        # ====================================================
        # OUTPUT COLLECTIONS FOR THIS MONTH
        # ====================================================

        transactions = []

        control_rows = []

        warnings = []

        legend_details = []

        color_to_label = {}


        # ====================================================
        # STRUCTURE VARIABLES
        # ====================================================

        header_row_number = None

        column_map = {}

        main_subtotal_row = None

        after_main_subtotal = False


        # ====================================================
        # STREAM ROWS
        # ========================================================
        #
        # KEY FIX:
        #
        # We use enumerate() for the row number.
        #
        # We NEVER use:
        #
        #     row[0].row
        #
        # because row[0] may be an EmptyCell.
        # ====================================================

        row_iterator = ws.iter_rows(

            min_row=1,

            max_col=MAX_COLUMNS_TO_READ,
        )


        for row_number, row in enumerate(

            row_iterator,

            start=1,
        ):


            # =================================================
            # BEFORE HEADER:
            # FIND LEGEND + TABLE HEADER
            # =================================================

            if header_row_number is None:


                normalized_values = [

                    normalize_text(
                        getattr(
                            cell,
                            "value",
                            None,
                        )
                    )

                    for cell
                    in row
                ]


                # =============================================
                # LEGACY ASCOT COLOR LEGEND
                # =============================================

                for column_number, cell in enumerate(

                    row,

                    start=1,
                ):


                    text = normalize_text(

                        getattr(
                            cell,
                            "value",
                            None,
                        )
                    )


                    if (
                        text
                        not in KNOWN_LEGEND_LABELS
                    ):

                        continue


                    presentation_label = (

                        KNOWN_LEGEND_LABELS[
                            text
                        ]
                    )


                    # First try the label cell itself.
                    fill_key = get_fill_key(
                        cell
                    )


                    # If the text cell itself has no color,
                    # inspect the immediately adjacent cells.
                    if fill_key is None:


                        neighbor_positions = [

                            column_number - 1,

                            column_number + 1,
                        ]


                        for neighbor_column in (
                            neighbor_positions
                        ):


                            if not (
                                1
                                <= neighbor_column
                                <= len(row)
                            ):

                                continue


                            neighbor_cell = row[
                                neighbor_column - 1
                            ]


                            neighbor_fill = (
                                get_fill_key(
                                    neighbor_cell
                                )
                            )


                            if neighbor_fill is not None:

                                fill_key = (
                                    neighbor_fill
                                )

                                break


                    signature = fill_signature(
                        fill_key
                    )


                    if signature:

                        color_to_label[
                            signature
                        ] = presentation_label


                    # KEY FIX:
                    #
                    # Do NOT use cell.column.
                    #
                    # column_number comes from enumerate(),
                    # so EmptyCell cannot cause an error.
                    legend_details.append(

                        {
                            "label":
                                presentation_label,

                            "source_row":
                                row_number,

                            "source_column":
                                column_number,

                            "fill":
                                fill_key,
                        }
                    )


                # =============================================
                # FIND ACTUAL AIC WF HEADER ROW
                # =============================================

                expected_anchor_headers = {

                    "as-of date",

                    "acct no",

                    "tran desc",

                    "debit amt",

                    "credit amt",
                }


                matches = (

                    expected_anchor_headers

                    & set(
                        normalized_values
                    )
                )


                if len(
                    matches
                ) >= 4:


                    header_row_number = (
                        row_number
                    )


                    # Map Excel positions to canonical fields.
                    for column_number, cell in enumerate(

                        row,

                        start=1,
                    ):


                        normalized_header = (
                            normalize_text(

                                getattr(
                                    cell,
                                    "value",
                                    None,
                                )
                            )
                        )


                        if (
                            normalized_header
                            in HEADER_MAP
                        ):

                            column_map[
                                column_number
                            ] = (
                                HEADER_MAP[
                                    normalized_header
                                ]
                            )


                    print(
                        f"Detected header row: "
                        f"{header_row_number}"
                    )


                    print(
                        f"Detected data columns: "
                        f"{len(column_map)}"
                    )


                continue


            # =================================================
            # BUILD ROW VALUES FROM CANONICAL COLUMNS
            # =================================================

            values = {}


            for column_number, field_name in (
                column_map.items()
            ):


                # Safety check in case a future workbook has
                # a shorter streamed tuple.
                if column_number > len(row):

                    value = None

                else:

                    cell = row[
                        column_number - 1
                    ]

                    value = getattr(
                        cell,
                        "value",
                        None,
                    )


                values[
                    field_name
                ] = value


            # =================================================
            # DETECT SUBTOTAL / CONTROL ROW
            # =================================================

            (
                has_sum_formula,
                formulas,
            ) = detect_sum_formula(
                row
            )


            # These are the fields that identify a normal
            # bank transaction.
            transaction_identity_present = any(

                values.get(
                    field_name
                )
                not in {
                    None,
                    "",
                }

                for field_name in [

                    "as_of_date",

                    "account_number",

                    "transaction_description",

                    "check_number",
                ]
            )


            if (
                has_sum_formula
                and not transaction_identity_present
            ):


                # First subtotal = main section boundary.
                if main_subtotal_row is None:

                    main_subtotal_row = (
                        row_number
                    )

                    after_main_subtotal = (
                        True
                    )

                    control_type = (
                        "main_cash_activity_subtotal"
                    )


                else:

                    control_type = (
                        "other_control_subtotal"
                    )


                control_rows.append(

                    {
                        "accounting_period":
                            period,

                        "processing_sequence":
                            processing_sequence,

                        "source_workbook":
                            workbook_path.name,

                        "source_sheet":
                            "AIC WF",

                        "source_row":
                            row_number,

                        "control_type":
                            control_type,

                        "formulas":
                            json.dumps(
                                formulas
                            ),
                    }
                )


                # Subtotal is not a transaction.
                continue


            # =================================================
            # DETERMINE WHETHER ROW IS A TRANSACTION
            # =================================================

            key_transaction_fields = [

                "as_of_date",

                "account_number",

                "transaction_description",

                "debit_amount",

                "credit_amount",

                "check_number",
            ]


            looks_like_transaction = any(

                values.get(
                    field_name
                )
                not in {
                    None,
                    "",
                }

                for field_name
                in key_transaction_fields
            )


            # Blank/spacer/presentation-only row.
            if not looks_like_transaction:

                continue


            # =================================================
            # NORMALIZE DATE
            # =================================================

            values[
                "as_of_date"
            ] = normalize_date(

                values.get(
                    "as_of_date"
                )
            )


            # =================================================
            # NORMALIZE NUMBERS
            # =================================================

            for numeric_field in [

                "debit_amount",

                "credit_amount",

                "zero_day_float_amount",
            ]:


                if numeric_field in values:

                    values[
                        numeric_field
                    ] = normalize_number(

                        values.get(
                            numeric_field
                        )
                    )


            # =================================================
            # IDENTIFY ZBA / SWEEP
            # =================================================

            zba_sweep = is_zba_or_sweep(

                values.get(
                    "transaction_description"
                ),

                values.get(
                    "descriptive_text_1"
                ),
            )


            # =================================================
            # ASSIGN LOGICAL SECTION
            # =================================================

            if not after_main_subtotal:

                section = (
                    "main_cash_activity"
                )

                section_order = 1


            elif zba_sweep:

                section = (
                    "zba_sweep_transfers"
                )

                section_order = 2


            else:

                # Preserve rather than silently force it into
                # the ZBA/sweep population.
                section = (
                    "post_subtotal_other"
                )

                section_order = 3


                warnings.append(

                    {
                        "accounting_period":
                            period,

                        "source_row":
                            row_number,

                        "transaction_description":
                            values.get(
                                "transaction_description"
                            ),

                        "warning":
                            (
                                "Transaction occurs after "
                                "main subtotal but is not "
                                "recognized as ZBA/sweep."
                            ),
                    }
                )


            # =================================================
            # READ LEGACY ASCOT COLOR INFORMATION
            # =================================================

            legacy_labels = []

            legacy_fill_signatures = []


            for column_number in (
                column_map.keys()
            ):


                if column_number > len(row):

                    continue


                cell = row[
                    column_number - 1
                ]


                signature = fill_signature(

                    get_fill_key(
                        cell
                    )
                )


                if signature is None:

                    continue


                if signature in (
                    color_to_label
                ):


                    presentation_label = (

                        color_to_label[
                            signature
                        ]
                    )


                    if (
                        presentation_label
                        not in legacy_labels
                    ):

                        legacy_labels.append(
                            presentation_label
                        )


                    if (
                        signature
                        not in legacy_fill_signatures
                    ):

                        legacy_fill_signatures.append(
                            signature
                        )


            # =================================================
            # DERIVE SIMPLE BANK-SIDE DIRECTION + AMOUNT
            # =================================================

            debit = values.get(
                "debit_amount"
            )

            credit = values.get(
                "credit_amount"
            )


            has_debit = (
                debit is not None
                and debit != 0
            )

            has_credit = (
                credit is not None
                and credit != 0
            )


            if (
                has_debit
                and not has_credit
            ):

                bank_direction = (
                    "debit"
                )

                bank_amount = abs(
                    debit
                )


            elif (
                has_credit
                and not has_debit
            ):

                bank_direction = (
                    "credit"
                )

                bank_amount = abs(
                    credit
                )


            elif (
                has_debit
                and has_credit
            ):

                bank_direction = (
                    "mixed"
                )

                bank_amount = None


            else:

                bank_direction = None

                bank_amount = None


            # =================================================
            # ADD ACCOUNTING PERIOD / CHRONOLOGY
            # =================================================

            values[
                "accounting_period"
            ] = period


            values[
                "processing_sequence"
            ] = processing_sequence


            # =================================================
            # ADD SOURCE PROVENANCE
            # =================================================

            values[
                "source_workbook"
            ] = workbook_path.name


            values[
                "source_sheet"
            ] = "AIC WF"


            values[
                "source_row"
            ] = row_number


            # =================================================
            # ADD MACHINE-READABLE STRUCTURE
            # =================================================

            values[
                "section"
            ] = section


            values[
                "section_order"
            ] = section_order


            values[
                "is_zba_or_sweep"
            ] = zba_sweep


            values[
                "included_in_core_cash_activity"
            ] = (
                section
                == "main_cash_activity"
            )


            values[
                "bank_direction"
            ] = bank_direction


            values[
                "bank_amount"
            ] = bank_amount


            # =================================================
            # PRESERVE LEGACY PRESENTATION
            # =================================================

            values[
                "legacy_presentation_labels"
            ] = (

                "|".join(
                    legacy_labels
                )

                if legacy_labels

                else None
            )


            values[
                "legacy_fill_signatures"
            ] = (

                "|".join(
                    legacy_fill_signatures
                )

                if legacy_fill_signatures

                else None
            )


            # =================================================
            # FUTURE MACHINE CATEGORY
            # =================================================
            #
            # This is intentionally separate from legacy color.
            #
            # Future deterministic / AI classification logic
            # will populate this field.
            # =================================================

            values[
                "machine_category"
            ] = None


            transactions.append(
                values
            )


        # ====================================================
        # VALIDATE HEADER DETECTION
        # ====================================================

        if header_row_number is None:

            raise ValueError(

                "Could not identify the AIC WF header row."
            )


        # ====================================================
        # BUILD SECTION POPULATIONS
        # ====================================================

        main_transactions = [

            tx

            for tx
            in transactions

            if tx[
                "section"
            ] == "main_cash_activity"
        ]


        zba_sweep_transactions = [

            tx

            for tx
            in transactions

            if tx[
                "section"
            ] == "zba_sweep_transfers"
        ]


        unexpected_post_subtotal = [

            tx

            for tx
            in transactions

            if tx[
                "section"
            ] == "post_subtotal_other"
        ]


        # ====================================================
        # CREATE MONTH-SPECIFIC OUTPUT FOLDER
        # ====================================================

        output_folder = (

            PROCESSED_FOLDER

            / workbook_path.stem
        )


        output_folder.mkdir(

            parents=True,

            exist_ok=True,
        )


        # ====================================================
        # WRITE NORMALIZED TRANSACTIONS
        # ====================================================

        write_csv(

            output_folder
            / "aic_wf_transactions.csv",

            transactions,
        )


        # ====================================================
        # WRITE CONTROL / SUBTOTAL ROWS
        # ====================================================

        write_csv(

            output_folder
            / "aic_wf_control_rows.csv",

            control_rows,
        )


        # ====================================================
        # WRITE PRESENTATION METADATA
        # ====================================================

        presentation_metadata = {

            "accounting_period":
                period,

            "processing_sequence":
                processing_sequence,

            "source_workbook":
                workbook_path.name,

            "source_sheet":
                "AIC WF",

            "header_row":
                header_row_number,

            "main_subtotal_row":
                main_subtotal_row,

            "color_legend":
                legend_details,

            "layout_rule": {

                "first_section":
                    "main_cash_activity",

                "boundary":
                    "main_cash_activity_subtotal",

                "second_section":
                    "zba_sweep_transfers",

                "preserve_visual_gap_on_export":
                    True,

                "do_not_use_gap_alone_to_identify_sweeps":
                    True,
            },
        }


        with (
            output_folder
            / "aic_wf_presentation.json"
        ).open(

            "w",

            encoding="utf-8",

        ) as file:


            json.dump(

                presentation_metadata,

                file,

                indent=2,

                ensure_ascii=False,
            )


        # ====================================================
        # WRITE VALIDATION SUMMARY
        # ====================================================

        validation_summary = {

            "accounting_period":
                period,

            "processing_sequence":
                processing_sequence,

            "source_workbook":
                workbook_path.name,

            "header_row":
                header_row_number,

            "main_subtotal_row":
                main_subtotal_row,

            "transaction_count":
                len(
                    transactions
                ),

            "main_cash_activity_count":
                len(
                    main_transactions
                ),

            "zba_sweep_count":
                len(
                    zba_sweep_transactions
                ),

            "unexpected_post_subtotal_count":
                len(
                    unexpected_post_subtotal
                ),

            "control_row_count":
                len(
                    control_rows
                ),

            "warning_count":
                len(
                    warnings
                ),

            "warnings":
                warnings,
        }


        with (
            output_folder
            / "aic_wf_validation.json"
        ).open(

            "w",

            encoding="utf-8",

        ) as file:


            json.dump(

                validation_summary,

                file,

                indent=2,

                ensure_ascii=False,
            )


        # ====================================================
        # TERMINAL SUMMARY
        # ====================================================

        print(
            f"Main subtotal row: "
            f"{main_subtotal_row}"
        )


        print(
            f"Main cash transactions: "
            f"{len(main_transactions)}"
        )


        print(
            f"ZBA/sweep transactions: "
            f"{len(zba_sweep_transactions)}"
        )


        print(
            "Unexpected post-subtotal transactions: "
            f"{len(unexpected_post_subtotal)}"
        )


        print(
            f"Control rows: "
            f"{len(control_rows)}"
        )


        print(
            f"Legend mappings detected: "
            f"{len(color_to_label)}"
        )


        if unexpected_post_subtotal:

            print(
                "⚠ Review unexpected "
                "post-subtotal transactions."
            )


        else:

            print(
                "✓ Post-subtotal structure follows "
                "the expected ZBA/sweep pattern."
            )


        return (
            transactions,
            control_rows,
            validation_summary,
        )


    finally:

        # Always close the workbook even if this month's
        # processing produces an exception.
        wb.close()


# ============================================================
# MAIN PROGRAM
# ============================================================

def main():

    print(
        "ASCOT AIC WF SANITIZER"
    )

    print(
        "Streaming + chronological + EmptyCell-safe mode"
    )


    # ========================================================
    # FIND RESOURCE PRO FILES
    # ========================================================

    files = list(

        RAW_FOLDER.glob(
            "*Resource Pro Cash.xlsx"
        )
    )


    if not files:

        raise FileNotFoundError(

            "No Resource Pro Cash files were found "
            "inside data/raw/."
        )


    # ========================================================
    # CHRONOLOGICAL SORT
    # ========================================================

    files = sorted(

        files,

        key=extract_accounting_period,
    )


    print(
        f"\nFound {len(files)} workbook(s)."
    )


    print(
        "\nChronological processing order:"
    )


    for sequence, file_path in enumerate(

        files,

        start=1,
    ):


        print(

            f"  {sequence}. "
            f"{file_path.name}"
        )


    # ========================================================
    # CONSOLIDATED OUTPUT COLLECTIONS
    # ========================================================

    all_transactions = []

    all_control_rows = []

    processing_manifest = []


    # ========================================================
    # PROCESS MONTHS IN ACCOUNTING-PERIOD ORDER
    # ========================================================

    for processing_sequence, workbook_path in enumerate(

        files,

        start=1,
    ):


        try:

            (
                transactions,
                control_rows,
                validation,
            ) = sanitize_workbook(

                workbook_path,

                processing_sequence,
            )


            all_transactions.extend(
                transactions
            )


            all_control_rows.extend(
                control_rows
            )


            processing_manifest.append(

                {
                    "processing_sequence":
                        processing_sequence,

                    "accounting_period":
                        accounting_period_label(
                            workbook_path
                        ),

                    "source_workbook":
                        workbook_path.name,

                    "status":
                        "success",

                    "transaction_count":
                        validation[
                            "transaction_count"
                        ],

                    "main_cash_activity_count":
                        validation[
                            "main_cash_activity_count"
                        ],

                    "zba_sweep_count":
                        validation[
                            "zba_sweep_count"
                        ],

                    "unexpected_post_subtotal_count":
                        validation[
                            "unexpected_post_subtotal_count"
                        ],
                }
            )


        except Exception as error:

            print(
                "\nERROR while processing "
                f"{workbook_path.name}:"
            )


            print(
                f"{type(error).__name__}: "
                f"{error}"
            )


            processing_manifest.append(

                {
                    "processing_sequence":
                        processing_sequence,

                    "accounting_period":
                        accounting_period_label(
                            workbook_path
                        ),

                    "source_workbook":
                        workbook_path.name,

                    "status":
                        "error",

                    "error_type":
                        type(error).__name__,

                    "error":
                        str(error),
                }
            )


    # ========================================================
    # EXPLICITLY SORT CONSOLIDATED OUTPUT CHRONOLOGICALLY
    # ========================================================
    #
    # This is intentionally done even though we already
    # processed the source workbooks chronologically.
    #
    # It makes chronology an explicit pipeline invariant.
    # ========================================================

    all_transactions.sort(

        key=lambda record: (

            record[
                "processing_sequence"
            ],

            record[
                "source_row"
            ],
        )
    )


    all_control_rows.sort(

        key=lambda record: (

            record[
                "processing_sequence"
            ],

            record[
                "source_row"
            ],
        )
    )


    # ========================================================
    # WRITE ALL-MONTH TRANSACTION TABLE
    # ========================================================

    write_csv(

        PROCESSED_FOLDER
        / "aic_wf_all_months_transactions.csv",

        all_transactions,
    )


    # ========================================================
    # WRITE ALL-MONTH CONTROL TABLE
    # ========================================================

    write_csv(

        PROCESSED_FOLDER
        / "aic_wf_all_months_control_rows.csv",

        all_control_rows,
    )


    # ========================================================
    # WRITE PROCESSING MANIFEST
    # ========================================================

    with (
        PROCESSED_FOLDER
        / "aic_wf_processing_manifest.json"
    ).open(

        "w",

        encoding="utf-8",

    ) as file:


        json.dump(

            processing_manifest,

            file,

            indent=2,

            ensure_ascii=False,
        )


    # ========================================================
    # FINAL SUMMARY
    # ========================================================

    successful_months = sum(

        item[
            "status"
        ] == "success"

        for item
        in processing_manifest
    )


    failed_months = (

        len(
            processing_manifest
        )

        - successful_months
    )


    print(
        "\n========================================"
    )

    print(
        "SANITIZATION COMPLETE"
    )

    print(
        "========================================"
    )


    print(
        f"Successful months: "
        f"{successful_months}"
    )


    print(
        f"Failed months: "
        f"{failed_months}"
    )


    print(
        f"Total normalized transactions: "
        f"{len(all_transactions)}"
    )


    print(
        f"Total control rows: "
        f"{len(all_control_rows)}"
    )


    print(
        "\nConsolidated transaction output:"
    )


    print(
        "data/processed/"
        "aic_wf_all_months_transactions.csv"
    )


    print(
        "\nProcessing manifest:"
    )


    print(
        "data/processed/"
        "aic_wf_processing_manifest.json"
    )


# ============================================================
# RUN PROGRAM
# ============================================================

if __name__ == "__main__":

    main()